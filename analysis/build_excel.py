"""Build Sales_Dashboard.xlsx: Cover + Data + Dashboard (KPIs, tables, charts)."""
import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

BASE = Path(__file__).resolve().parent.parent
df = pd.read_csv(BASE / "data" / "sales_data_clean.csv", parse_dates=["date"])
monthly = pd.read_csv(BASE / "data" / "monthly_summary.csv", parse_dates=["month"])

NAVY, TEAL, LIGHT = "1F3A5F", "2E86AB", "D6EAF8"
title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
head_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
head_fill = PatternFill("solid", fgColor=NAVY)
kpi_fill = PatternFill("solid", fgColor=LIGHT)
thin = Border(*[Side(style="thin", color="B0B0B0")] * 4)
center = Alignment(horizontal="center", vertical="center")

wb = Workbook()

def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = head_font; cell.fill = head_fill
        cell.alignment = center; cell.border = thin

# ---------- Cover ----------
ws = wb.active; ws.title = "Cover"
ws.sheet_properties.tabColor = NAVY
ws["A1"] = "Dar El Baraka - Sales Dashboard"
ws["A1"].font = Font(name="Calibri", size=20, bold=True, color=NAVY)
ws["A3"] = ("Fictional Tunisian SME distributing local food products (olive oil, dates, "
            "harissa, grains, honey). 18 months of orders: Apr 2024 - Sep 2025.")
ws["A3"].alignment = Alignment(wrap_text=True); ws["A3"].font = Font(size=11)
ws.column_dimensions["A"].width = 110
ws["A5"] = "Sheets:"; ws["A5"].font = Font(bold=True, size=12)
for i, s in enumerate(["Data - cleaned order lines (2,460 rows)",
                       "Dashboard - KPIs, summary tables and charts",
                       "Monthly - month-by-month revenue/profit/orders"], start=6):
    ws[f"A{i}"] = f"- {s}"
ws["A10"] = "Built with Python (pandas + matplotlib + openpyxl). See README.md for the full case study."

# ---------- Data ----------
ws = wb.create_sheet("Data")
ws.sheet_properties.tabColor = TEAL
for c, col in enumerate(df.columns, start=1):
    ws.cell(row=1, column=c, value=col)
style_header(ws, 1, len(df.columns))
for r, row in enumerate(df.itertuples(index=False), start=2):
    for c, val in enumerate(row, start=1):
        cell = ws.cell(row=r, column=c, value=val)
        cell.border = thin
        if isinstance(val, pd.Timestamp):
            cell.number_format = "YYYY-MM-DD"
for c in range(1, len(df.columns) + 1):
    ws.column_dimensions[get_column_letter(c)].width = 18
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{len(df)+1}"

# ---------- Dashboard ----------
ws = wb.create_sheet("Dashboard")
ws.sheet_properties.tabColor = "27AE60"
ws["A1"] = "KEY PERFORMANCE INDICATORS"; ws["A1"].font = title_font
ws["A1"].fill = PatternFill("solid", fgColor=NAVY); ws["A1"].alignment = center
ws.merge_cells("A1:D1")

total_rev = df["revenue_tnd"].sum(); total_profit = df["profit_tnd"].sum()
n_orders = df["order_id"].nunique()
kpis = [("Total Revenue (TND)", f"{total_rev:,.0f}"),
        ("Total Profit (TND)", f"{total_profit:,.0f}"),
        ("Profit Margin", f"{total_profit/total_rev*100:.1f}%"),
        ("Orders", f"{n_orders:,}"),
        ("Avg Order Value (TND)", f"{total_rev/n_orders:,.1f}"),
        ("Period", "Apr 2024 - Sep 2025")]
for i, (k, v) in enumerate(kpis, start=2):
    ws.cell(row=i, column=1, value=k).font = Font(bold=True)
    c = ws.cell(row=i, column=2, value=v); c.font = Font(bold=True, size=12, color=NAVY)
    for col in (1, 2):
        ws.cell(row=i, column=col).fill = kpi_fill; ws.cell(row=i, column=col).border = thin
ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 22

r0 = 10
ws[f"A{r0}"] = "REVENUE & MARGIN BY CHANNEL"; ws[f"A{r0}"].font = title_font
ws[f"A{r0}"].fill = PatternFill("solid", fgColor=NAVY); ws.merge_cells(f"A{r0}:D{r0}")
ch = df.groupby("channel").agg(revenue=("revenue_tnd", "sum"), profit=("profit_tnd", "sum"),
                               orders=("order_id", "nunique"))
ch["margin %"] = (ch["profit"] / ch["revenue"] * 100).round(1)
hdr = ["Channel", "Revenue (TND)", "Profit (TND)", "Margin %"]
for c, h in enumerate(hdr, start=1):
    ws.cell(row=r0 + 1, column=c, value=h)
style_header(ws, r0 + 1, 4)
for i, (name, row) in enumerate(ch.iterrows(), start=r0 + 2):
    ws.cell(row=i, column=1, value=name).font = Font(bold=True)
    ws.cell(row=i, column=2, value=round(row["revenue"])).number_format = "#,##0"
    ws.cell(row=i, column=3, value=round(row["profit"])).number_format = "#,##0"
    ws.cell(row=i, column=4, value=row["margin %"] / 100).number_format = "0.0%"
    for c in range(1, 5):
        ws.cell(row=i, column=c).border = thin

r1 = r0 + 7
ws[f"A{r1}"] = "MARGIN VS DISCOUNT BY CATEGORY"; ws[f"A{r1}"].font = title_font
ws[f"A{r1}"].fill = PatternFill("solid", fgColor=NAVY); ws.merge_cells(f"A{r1}:D{r1}")
cat = df.groupby("category").agg(revenue=("revenue_tnd", "sum"), profit=("profit_tnd", "sum"),
                                 avg_discount=("discount_pct", "mean"))
cat["margin %"] = (cat["profit"] / cat["revenue"] * 100).round(1)
for c, h in enumerate(["Category", "Revenue (TND)", "Margin %", "Avg Discount %"], start=1):
    ws.cell(row=r1 + 1, column=c, value=h)
style_header(ws, r1 + 1, 4)
for i, (name, row) in enumerate(cat.iterrows(), start=r1 + 2):
    ws.cell(row=i, column=1, value=name).font = Font(bold=True)
    ws.cell(row=i, column=2, value=round(row["revenue"])).number_format = "#,##0"
    ws.cell(row=i, column=3, value=row["margin %"] / 100).number_format = "0.0%"
    ws.cell(row=i, column=4, value=round(row["avg_discount"], 1) / 100).number_format = "0.0%"
    for c in range(1, 5):
        ws.cell(row=i, column=c).border = thin

# charts
img_row = r1 + 10
for j, img_file in enumerate(["monthly_revenue_by_channel.png", "top_products.png",
                              "category_margin_discount.png", "revenue_by_region.png"]):
    img = XLImage(str(BASE / "visuals" / img_file))
    img.width, img.height = 520, 234
    ws.add_image(img, f"A{img_row + j * 14}")

ws.column_dimensions["C"].width = 20; ws.column_dimensions["D"].width = 20

# ---------- Monthly ----------
ws = wb.create_sheet("Monthly")
for c, col in enumerate(["month", "revenue_tnd", "profit_tnd", "orders", "mom_growth_%"], start=1):
    ws.cell(row=1, column=c, value=col)
style_header(ws, 1, 5)
for i, row in enumerate(monthly.itertuples(), start=2):
    ws.cell(row=i, column=1, value=row.month).number_format = "YYYY-MM"
    ws.cell(row=i, column=2, value=round(row.revenue)).number_format = "#,##0"
    ws.cell(row=i, column=3, value=round(row.profit)).number_format = "#,##0"
    ws.cell(row=i, column=4, value=row.orders)
    ws.cell(row=i, column=5,
            value=(row.mom_growth / 100 if pd.notna(row.mom_growth) else None)).number_format = "0.0%"
    for c in range(1, 6):
        ws.cell(row=i, column=c).border = thin
for c in range(1, 6):
    ws.column_dimensions[get_column_letter(c)].width = 18

out = BASE / "dashboard" / "Sales_Dashboard.xlsx"
wb.save(out)
print("saved", out)
